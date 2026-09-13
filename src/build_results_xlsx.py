"""Consolidated results workbook — every metric, every model, reviewable."""
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import config as C

R = C.RESULT_DIR
agg = json.loads((R / "aggregated.json").read_text())
mm = json.loads((R / "mixed_models.json").read_text())
ex = json.loads((R / "extra_analyses.json").read_text())
op = json.loads((R / "operating_points.json").read_text())
gl = json.loads((R / "gain_lift.json").read_text())

ORDER = ["det_fusion", "det_count", "det_count_sp", "ens_cnn", "yolo_tl", "yolo_coco",
         "efficientnet_b3", "resnet50v2", "yolo_scratch"]
ORDER = [m for m in ORDER if m in agg]
NAME = C.MODEL_DISPLAY

A = lambda **k: Font(name="Arial", **k)
HF = PatternFill("solid", fgColor="1F4E79")
PROP = PatternFill("solid", fgColor="FFF2CC")
_t = Side("thin", color="BFBFBF"); BD = Border(_t, _t, _t, _t)


def sheet(wb, title):
    ws = wb.create_sheet(title) if wb.sheetnames != ["Sheet"] else wb.active
    if ws.title == "Sheet":
        ws.title = title
    return ws


def header(ws, cols, row=1):
    for j, c in enumerate(cols, 1):
        x = ws.cell(row, j, c)
        x.font = A(bold=True, color="FFFFFF"); x.fill = HF
        x.alignment = Alignment(wrap_text=True, vertical="center"); x.border = BD


def body(ws, rows, start=2, prop_col=0):
    for i, r in enumerate(rows):
        for j, v in enumerate(r, 1):
            c = ws.cell(start + i, j, v)
            c.font = A(size=10); c.border = BD
            c.alignment = Alignment(vertical="center", horizontal="left" if j == 1 else "center")
        if prop_col and r[0] and "proposed" in str(r[0]).lower():
            for j in range(1, len(r) + 1):
                ws.cell(start + i, j).fill = PROP


wb = Workbook()

# ---- Sheet 1: headline (mean ± sd, %) ----
ws = sheet(wb, "1 - Headline metrics")
METR = [("f1_macro", "macro-F1"), ("roc_auc", "ROC-AUC"), ("pr_auc_non_urgent", "PR-AUC non-urgent"),
        ("balanced_accuracy", "balanced acc"), ("mcc", "MCC"), ("accuracy", "accuracy"),
        ("f1_urgent", "F1 urgent"), ("f1_non_urgent", "F1 non-urgent"),
        ("precision_urgent", "prec urgent"), ("recall_urgent", "recall urgent"),
        ("precision_non_urgent", "prec non-urg"), ("recall_non_urgent", "recall non-urg")]
header(ws, ["model", "macro-F1 95% CI"] + [f"{lbl}\n(mean ± sd, %)" for _, lbl in METR])
rows = []
for m in ORDER:
    a = agg[m]["aggregate"]
    ci = a["f1_macro"]["ci95"]
    row = [NAME[m], f"[{ci[0]*100:.1f}, {ci[1]*100:.1f}]"]
    for key, _ in METR:
        d = a[key]
        row.append(f"{d['mean']*100:.1f} ± {d['sd']*100:.1f}")
    rows.append(row)
body(ws, rows, prop_col=1)
ws.freeze_panes = "B2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(METR)+2)}{len(rows)+1}"
for j, w in enumerate([44, 18] + [15] * len(METR), 1):
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[1].height = 42

# ---- Sheet 2: per-fold ----
ws2 = sheet(wb, "2 - Per-fold values")
for mi, (key, lbl) in enumerate([("f1_macro", "macro-F1"), ("roc_auc", "ROC-AUC"),
                                 ("f1_non_urgent", "F1 non-urgent"), ("mcc", "MCC")]):
    base = mi * (len(ORDER) + 3) + 1
    ws2.cell(base, 1, lbl).font = A(bold=True, size=11)
    header(ws2, ["model", "fold 1", "fold 2", "fold 3", "fold 4", "fold 5", "mean", "sd"], row=base + 1)
    rr = []
    for m in ORDER:
        vals = agg[m]["aggregate"][key]["values"]
        rr.append([NAME[m]] + [round(v * 100, 1) for v in vals] +
                  [round(np.mean(vals) * 100, 1), round(np.std(vals, ddof=1) * 100, 1)])
    body(ws2, rr, start=base + 2, prop_col=1)
for j, w in enumerate([44, 9, 9, 9, 9, 9, 9, 9], 1):
    ws2.column_dimensions[get_column_letter(j)].width = w

# ---- Sheet 3: statistical comparison vs proposed ----
ws3 = sheet(wb, "3 - Stats vs proposed")
ws3.cell(1, 1, f"Reference (proposed) = {NAME['det_fusion']}. "
               "GLMM: odds ratio <1 means that model is less likely to be correct than the proposed. "
               "Cluster bootstrap: pooled-OOF AUC difference (proposed - model), folds then patients resampled. "
               "Bayesian: P(proposed better) on per-fold macro-F1 / ROC-AUC.").font = A(size=9, italic=True)
ws3.merge_cells("A1:H1"); ws3.row_dimensions[1].height = 42
header(ws3, ["model", "GLMM odds ratio [95% CI]", "GLMM p",
             "Δ ROC-AUC [95% CI], p", "Δ PR-AUC non-urg [95% CI], p",
             "Bayesian P(prop. better)\nmacro-F1 / ROC-AUC",
             "paired-t p (macro-F1)", "NB-corrected sd (macro-F1)"], row=3)
gl_ct = {k: v for k, v in mm["glmm"].get("contrasts", {}).items()}
cb = mm["cluster_bootstrap"]; nb = mm["nadeau_bengio_macro_f1"]
bt = ex["bayesian_correlated_t"]
sc = json.loads((R / "stats_comparison.json").read_text())["per_fold_tests"]
rr = []
for m in ORDER:
    if m == "det_fusion":
        continue
    g = gl_ct.get(m, {})
    c = cb.get(m, {})
    b = bt.get(m, {})
    rr.append([
        NAME[m],
        f"{g.get('odds_ratio',float('nan')):.2f} [{g.get('or_ci95',[float('nan')]*2)[0]:.2f}, {g.get('or_ci95',[float('nan')]*2)[1]:.2f}]" if g else "-",
        f"{g.get('p_value',float('nan')):.3f}" if g else "-",
        (f"{c['roc_auc']['delta']*100:+.1f} [{c['roc_auc']['ci95'][0]*100:+.1f}, {c['roc_auc']['ci95'][1]*100:+.1f}], p={c['roc_auc']['p_cluster_bootstrap']:.3f}" if c else "-"),
        (f"{c['pr_auc_non_urgent']['delta']*100:+.1f} [{c['pr_auc_non_urgent']['ci95'][0]*100:+.1f}, {c['pr_auc_non_urgent']['ci95'][1]*100:+.1f}], p={c['pr_auc_non_urgent']['p_cluster_bootstrap']:.3f}" if c else "-"),
        (f"{b['f1_macro']['p_proposed_better']:.2f} / {b['roc_auc']['p_proposed_better']:.2f}" if b else "-"),
        f"{sc.get(m,{}).get('f1_macro',{}).get('paired_t_p',float('nan')):.3f}",
        f"{nb.get(m,{}).get('nadeau_bengio_sd',float('nan'))*100:.1f}",
    ])
body(ws3, rr, start=4)
for j, w in enumerate([44, 24, 9, 34, 34, 20, 14, 16], 1):
    ws3.column_dimensions[get_column_letter(j)].width = w
ws3.cell(4 + len(rr) + 1, 1,
         "GLMM variance components: patient_id "
         f"{mm['glmm'].get('variance_components',[{}])[0].get('vcov','?')}, fold "
         f"{mm['glmm'].get('variance_components',[{},{}])[-1].get('vcov','?')}.").font = A(size=9, italic=True)

# ---- Sheet 4: operating points ----
ws4 = sheet(wb, "4 - Operating points")
header(ws4, ["model", "operating point", "urgent sensitivity", "specificity",
             "NPV (defer safety)", "PPV", "cleared fraction", "missed urgent (sum/211)"])
rr = []
for m in ORDER:
    if m not in op:
        continue
    for pt in ("default_0.5", "sens>=90", "sens>=95"):
        a = op[m][pt]
        rr.append([NAME[m], pt,
                   f"{a['sensitivity']['mean']:.2f} ± {a['sensitivity']['sd']:.2f}",
                   f"{a['specificity']['mean']:.2f} ± {a['specificity']['sd']:.2f}",
                   f"{a['npv']['mean']:.2f} ± {a['npv']['sd']:.2f}",
                   f"{a['ppv']['mean']:.2f} ± {a['ppv']['sd']:.2f}",
                   f"{a['cleared_fraction']['mean']:.0%}",
                   f"{a['missed_urgent_total']}/{a['urgent_total']}"])
body(ws4, rr, prop_col=1)
for j, w in enumerate([44, 14, 18, 14, 18, 10, 14, 20], 1):
    ws4.column_dimensions[get_column_letter(j)].width = w

# ---- Sheet 5: calibration, gain/lift, trivial, subgroup ----
ws5 = sheet(wb, "5 - Calibration & context")
ws5.cell(1, 1, "Calibration (pooled OOF)").font = A(bold=True, size=11)
header(ws5, ["model", "Brier (lower better)", "ECE (lower better)"], row=2)
body(ws5, [[NAME[m], round(ex["calibration"][m]["brier"], 3), round(ex["calibration"][m]["ece"], 3)]
           for m in ORDER if m in ex["calibration"]], start=3, prop_col=1)
b = 3 + len(ORDER) + 2
ws5.cell(b, 1, "Worklist gain / lift (pooled OOF, model-ordered)").font = A(bold=True, size=11)
header(ws5, ["model", "urgent found in top 10% (lift)", "top 30% (lift)", "N reviewed to find 80% urgent"], row=b + 1)
rr = []
for m in ORDER:
    if m not in gl:
        continue
    d = gl[m]["at_deciles"]; nnr = gl[m]["number_needed_to_review"]
    rr.append([NAME[m],
               f"{d['top_10pct']['urgent_found']:.0%} ({d['top_10pct']['lift']:.2f}x)",
               f"{d['top_30pct']['urgent_found']:.0%} ({d['top_30pct']['lift']:.2f}x)",
               f"{nnr['to_find_80pct_urgent']['n_reviewed']} ({nnr['to_find_80pct_urgent']['pct_reviewed']:.0%})"])
body(ws5, rr, start=b + 2, prop_col=1)
b2 = b + 2 + len(rr) + 2
ws5.cell(b2, 1, "Trivial baselines (context)").font = A(bold=True, size=11)
header(ws5, ["strategy", "macro-F1", "F1 non-urgent", "accuracy"], row=b2 + 1)
body(ws5, [[k, round(v["f1_macro"] * 100, 1), round(v["f1_non_urgent"] * 100, 1), round(v["accuracy"] * 100, 1)]
           for k, v in ex["trivial_baselines"].items()], start=b2 + 2)
b3 = b2 + 2 + len(ex["trivial_baselines"]) + 2
ws5.cell(b3, 1, "Subgroup by source dataset (pooled OOF)").font = A(bold=True, size=11)
header(ws5, ["model", "source", "n", "macro-F1", "ROC-AUC"], row=b3 + 1)
rr = []
for m in ORDER:
    if m not in ex["subgroup_by_source"]:
        continue
    for s, d in ex["subgroup_by_source"][m].items():
        rr.append([NAME[m], s, d["n"], round(d["f1_macro"] * 100, 1),
                   round(d["roc_auc"] * 100, 1) if d["roc_auc"] else "-"])
body(ws5, rr, start=b3 + 2, prop_col=1)
for j, w in enumerate([44, 26, 20, 24], 1):
    ws5.column_dimensions[get_column_letter(j)].width = w

# ---- Sheet 0: notes ----
ws0 = wb.create_sheet("0 - Notes", 0)
for i, t in enumerate([
    "Pediatric dental case-level urgency (urgent vs non-urgent) from panoramic radiographs.",
    "287 unique patients, 211 urgent / 76 non-urgent (2.78:1). Labels = 3-rater consensus (Cohen's kappa 0.88).",
    "Protocol: patient-level nested 5-fold CV; Optuna (15 trials, macro-F1) on an inner 80/20 hold-out of each",
    "outer fold's own training data; each outer test fold scored once. Mirror/flip = train-only augmentation.",
    "",
    "Proposed = 'Detect-then-count + CNN (late fusion)': equal-weight (0.5) mean of P(urgent) from",
    "  (i) a tuned classifier on ~12 features engineered from a caries detector's outputs, and",
    "  (ii) a CNN on multi-scale features of the same caries-adapted backbone.",
    "The caries detector (best.pt) was trained on adult DENTEX only - 0 image overlap with the 287 (verified,",
    "exact + perceptual hash) - and never saw any urgency label. Detection features are label- and fold-independent.",
    "",
    "All values are mean +- sd over the 5 outer test folds unless noted. Proposed row highlighted.",
    "Generated from results/*.json.",
], 1):
    ws0.cell(i, 1, t).font = A(size=10, bold=(i == 1))
ws0.column_dimensions["A"].width = 110

out = C.REPORT_DIR / "results_all_models.xlsx"
wb.save(out)
print("saved", out, "| sheets:", wb.sheetnames)
